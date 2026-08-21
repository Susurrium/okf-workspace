"""Build the deterministic DOCX conversion fixture."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path(__file__).with_name("simple.docx")
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)


def _font(run: object, name: str, size: float, *, bold: bool = False) -> None:
    run.font.name = name  # type: ignore[attr-defined]
    run.font.size = Pt(size)  # type: ignore[attr-defined]
    run.font.bold = bold  # type: ignore[attr-defined]
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()  # type: ignore[attr-defined]
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)


def _hyperlink(paragraph: object, label: str, url: str) -> None:
    part = paragraph.part  # type: ignore[attr-defined]
    relationship = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    text = OxmlElement("w:t")
    text.text = label
    run.extend((properties, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)  # type: ignore[attr-defined]


def _set_cell_margins(cell: object) -> None:
    properties = cell._tc.get_or_add_tcPr()  # type: ignore[attr-defined]
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_table_geometry(table: object, widths: tuple[int, ...]) -> None:
    table.autofit = False  # type: ignore[attr-defined]
    properties = table._tbl.tblPr  # type: ignore[attr-defined]
    table_width = properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    properties.append(indent)
    grid = table._tbl.tblGrid  # type: ignore[attr-defined]
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:  # type: ignore[attr-defined]
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def build() -> Path:
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(10)
    title_run = title.add_run("OKF Document Conversion Fixture")
    _font(title_run, "Calibri", 24, bold=True)
    title_run.font.color.rgb = DARK_BLUE

    lead = document.add_paragraph(
        "This compact reference checks headings, real lists, hyperlinks, and fixed-width tables."
    )
    lead.paragraph_format.space_after = Pt(8)

    document.add_heading("Import checklist", level=1)
    for value in (
        "Preserve the complete body instead of a short preview.",
        "Keep the Original outside the synchronized Vault.",
        "Publish a Source Record and Source Text after conversion.",
    ):
        paragraph = document.add_paragraph(value, style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25

    link = document.add_paragraph("Project format reference: ")
    _hyperlink(link, "OKF specification", "https://openknowledge.fyi/")

    document.add_heading("Parser expectations", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ("Element", "Expected Markdown", "Acceptance")
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            _font(run, "Calibri", 10, bold=True)
    for values in (
        ("Heading", "## Import checklist", "Hierarchy retained"),
        ("List", "Three bullet items", "Items remain separate"),
        ("Table", "Three columns", "Rows remain readable"),
    ):
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
            for run in cell.paragraphs[0].runs:
                _font(run, "Calibri", 10)
    _set_table_geometry(table, (1800, 3600, 3960))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("OKF Brain - deterministic fixture")
    _font(run, "Calibri", 9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    document.core_properties.title = "OKF Document Conversion Fixture"
    document.core_properties.subject = "Deterministic MarkItDown conversion fixture"
    document.core_properties.author = "OKF Brain"
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
