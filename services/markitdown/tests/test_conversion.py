from __future__ import annotations

import base64
import hashlib
import io
import json
import zipfile
from pathlib import Path

from pypdf import PdfWriter

from markitdown_service.conversion import (
    ConversionLimitError,
    convert_document_to_bundle,
    convert_pdf_to_bundle,
)


def make_pdf(path: Path, pages: int = 1) -> None:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=72, height=72)
    with path.open("wb") as handle:
        writer.write(handle)


PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def make_image_docx(path: Path) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
            'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
            '<Default Extension="png" ContentType="image/png"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        )
        archive.writestr(
            "word/_rels/document.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rIdImage" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/image" Target="media/image.png"/></Relationships>',
        )
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            "<w:body><w:p><w:r><w:t>Illustrated document</w:t></w:r></w:p><w:p><w:r><w:drawing>"
            '<wp:inline><wp:extent cx="9525" cy="9525"/><wp:docPr id="1" name="Figure"/>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            '<pic:pic><pic:nvPicPr><pic:cNvPr id="1" name="image.png"/><pic:cNvPicPr/>'
            '</pic:nvPicPr><pic:blipFill><a:blip r:embed="rIdImage"/><a:stretch><a:fillRect/>'
            '</a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/>'
            '<a:ext cx="9525" cy="9525"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/>'
            "</a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:inline>"
            "</w:drawing></w:r></w:p></w:body></w:document>",
        )
        archive.writestr("word/media/image.png", PNG)


def test_conversion_builds_exact_verified_bundle(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    output = tmp_path / "result.zip"
    make_pdf(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    manifest = convert_pdf_to_bundle(
        source,
        output,
        source_sha256=source_hash,
        filename="source.pdf",
        profile="fallback-text-v1",
        max_pages=10,
        max_result_bytes=4096,
        engine_version="0.1.6",
        engine_build_id="markitdown-0.1.6-py31213-assets-v3",
        converter=lambda _handle: "# Test\n\nBody\n",
    )

    with zipfile.ZipFile(output) as archive:
        assert set(archive.namelist()) == {"manifest.json", "document.md"}
        markdown = archive.read("document.md")
        archived_manifest = json.loads(archive.read("manifest.json"))
        assert archive.getinfo("manifest.json").date_time == (1980, 1, 1, 0, 0, 0)

    assert manifest == archived_manifest
    assert manifest["schema"] == "markitdown-result/v3"
    assert manifest["source_sha256"] == source_hash
    assert manifest["input"]["unit_count"] == 1
    assert manifest["engine"]["build_id"] == "markitdown-0.1.6-py31213-assets-v3"
    assert manifest["capabilities"] == {
        "page_map": False,
        "assets": False,
        "ocr": False,
        "formulas": False,
        "tables": "best_effort",
        "headings": "best_effort",
    }
    assert manifest["input"]["size_bytes"] == source.stat().st_size
    assert manifest["input"]["source_format"] == "pdf"
    assert manifest["input"]["unit_kind"] == "page"
    assert manifest["input"]["unit_count"] == 1
    assert manifest["output"]["markdown_sha256"] == hashlib.sha256(markdown).hexdigest()


def test_pinned_markitdown_pdf_converter_runs_without_a_mock(tmp_path: Path) -> None:
    source = tmp_path / "real-converter.pdf"
    output = tmp_path / "real-converter.zip"
    make_pdf(source)

    manifest = convert_pdf_to_bundle(
        source,
        output,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        filename=source.name,
        profile="fallback-text-v1",
        max_pages=1,
        max_result_bytes=1024 * 1024,
        engine_version="0.1.6",
        engine_build_id="markitdown-0.1.6-py31213-assets-v3",
    )

    with zipfile.ZipFile(output) as archive:
        markdown = archive.read("document.md")
    assert markdown.decode("utf-8") == ""
    assert manifest["warnings"] == [
        "pdf_text_only_images_not_extracted",
        "empty_markdown",
    ]
    assert manifest["output"]["markdown_sha256"] == hashlib.sha256(markdown).hexdigest()


def test_page_limit_is_enforced_before_conversion(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    make_pdf(source, pages=2)
    called = False

    def converter(_handle: object) -> str:
        nonlocal called
        called = True
        return "should not run"

    try:
        convert_pdf_to_bundle(
            source,
            tmp_path / "result.zip",
            source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
            filename="source.pdf",
            profile="fallback-text-v1",
            max_pages=1,
            max_result_bytes=4096,
            engine_version="0.1.6",
            engine_build_id="markitdown-0.1.6-py31213-assets-v3",
            converter=converter,
        )
    except ConversionLimitError:
        pass
    else:
        raise AssertionError("expected the page limit to fail")
    assert called is False


def test_result_limit_does_not_publish_partial_zip(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    output = tmp_path / "result.zip"
    make_pdf(source)

    try:
        convert_pdf_to_bundle(
            source,
            output,
            source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
            filename="source.pdf",
            profile="fallback-text-v1",
            max_pages=1,
            max_result_bytes=8,
            engine_version="0.1.6",
            engine_build_id="markitdown-0.1.6-py31213-assets-v3",
            converter=lambda _handle: "too much Markdown",
        )
    except ConversionLimitError:
        pass
    else:
        raise AssertionError("expected the result limit to fail")
    assert not output.exists()


def test_pinned_markitdown_pptx_converter_preserves_slide_text(tmp_path: Path) -> None:
    from pptx import Presentation

    source = tmp_path / "slides.pptx"
    output = tmp_path / "slides.zip"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "课程标题"
    slide.placeholders[1].text = "第一条知识点"
    presentation.save(source)

    manifest = convert_document_to_bundle(
        source,
        output,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        filename=source.name,
        profile="fallback-text-v1",
        max_units=10,
        max_result_bytes=1024 * 1024,
        engine_version="0.1.6",
        engine_build_id="markitdown-0.1.6-py31213-assets-v3",
    )

    with zipfile.ZipFile(output) as archive:
        markdown = archive.read("document.md").decode("utf-8")
    assert manifest["input"]["source_format"] == "pptx"
    assert manifest["input"]["unit_kind"] == "slide"
    assert manifest["input"]["unit_count"] == 1
    assert "课程标题" in markdown
    assert "第一条知识点" in markdown


def test_pinned_markitdown_docx_converter_preserves_word_text(tmp_path: Path) -> None:
    source = tmp_path / "notes.docx"
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
            'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        )
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p><w:r><w:t>Word heading</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>Native Word body text.</w:t></w:r></w:p>"
            "</w:body></w:document>",
        )
    output = tmp_path / "result.zip"
    manifest = convert_document_to_bundle(
        source,
        output,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        filename=source.name,
        profile="fallback-text-v1",
        max_units=500,
        max_result_bytes=1024 * 1024,
        engine_version="0.1.6",
        engine_build_id="test-build",
    )
    with zipfile.ZipFile(output) as archive:
        markdown = archive.read("document.md").decode()
    assert "Word heading" in markdown
    assert "Native Word body text." in markdown
    assert manifest["input"]["size_bytes"] == source.stat().st_size
    assert manifest["input"]["source_format"] == "docx"
    assert manifest["input"]["unit_kind"] == "document"
    assert manifest["input"]["unit_count"] == 1


def test_docx_embedded_image_is_hashed_rewritten_and_manifested(tmp_path: Path) -> None:
    source = tmp_path / "illustrated.docx"
    output = tmp_path / "result.zip"
    make_image_docx(source)

    manifest = convert_document_to_bundle(
        source,
        output,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        filename=source.name,
        profile="fallback-assets-v2",
        max_units=500,
        max_result_bytes=1024 * 1024,
        engine_version="0.1.6",
        engine_build_id="test-build",
    )

    digest = hashlib.sha256(PNG).hexdigest()
    asset_path = f"images/{digest}.png"
    with zipfile.ZipFile(output) as archive:
        markdown = archive.read("document.md").decode("utf-8")
        assert archive.read(asset_path) == PNG
    assert f"]({asset_path})" in markdown
    assert "data:image/" not in markdown
    assert manifest["input"]["embedded_image_relationships"] == 1
    assert manifest["input"]["embedded_image_references"] == 1
    assert manifest["output"]["image_references"] == 1
    assert manifest["output"]["assets"] == [
        {
            "path": asset_path,
            "sha256": digest,
            "mime_type": "image/png",
            "size_bytes": len(PNG),
        }
    ]


def test_pptx_repeated_image_keeps_two_references_but_one_asset(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "repeated.pptx"
    output = tmp_path / "result.zip"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(PNG), 0, 0, width=Inches(1))
    slide.shapes.add_picture(io.BytesIO(PNG), Inches(2), 0, width=Inches(1))
    presentation.save(source)

    manifest = convert_document_to_bundle(
        source,
        output,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        filename=source.name,
        profile="fallback-assets-v2",
        max_units=20,
        max_result_bytes=1024 * 1024,
        engine_version="0.1.6",
        engine_build_id="test-build",
    )

    digest = hashlib.sha256(PNG).hexdigest()
    asset_path = f"images/{digest}.png"
    with zipfile.ZipFile(output) as archive:
        markdown = archive.read("document.md").decode("utf-8")
        assert archive.namelist().count(asset_path) == 1
    assert markdown.count(asset_path) == 2
    assert manifest["input"]["embedded_image_references"] == 2
    assert manifest["output"]["image_references"] == 2
    assert len(manifest["output"]["assets"]) == 1


def test_missing_ooxml_image_reference_is_reported_in_manifest(tmp_path: Path) -> None:
    source = tmp_path / "illustrated.docx"
    output = tmp_path / "result.zip"
    make_image_docx(source)

    manifest = convert_document_to_bundle(
        source,
        output,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        filename=source.name,
        profile="fallback-assets-v2",
        max_units=500,
        max_result_bytes=1024 * 1024,
        engine_version="0.1.6",
        engine_build_id="test-build",
        converter=lambda _handle: "# Text only\n\nThe converter omitted the image.\n",
    )

    assert manifest["input"]["embedded_image_references"] == 1
    assert manifest["output"]["image_references"] == 0
    assert "missing_image_references:1" in manifest["warnings"]


def test_svg_data_uri_is_not_published_and_becomes_a_missing_image_warning(
    tmp_path: Path,
) -> None:
    source = tmp_path / "unsafe-svg.docx"
    output = tmp_path / "result.zip"
    make_image_docx(source)
    malicious_svg = b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
    encoded = base64.b64encode(malicious_svg).decode("ascii")

    manifest = convert_document_to_bundle(
        source,
        output,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        filename=source.name,
        profile="fallback-assets-v2",
        max_units=500,
        max_result_bytes=1024 * 1024,
        engine_version="0.1.6",
        engine_build_id="test-build",
        converter=lambda _handle: f"# Unsafe\n\n![figure](data:image/svg+xml;base64,{encoded})\n",
    )

    with zipfile.ZipFile(output) as archive:
        markdown = archive.read("document.md").decode("utf-8")
        assert set(archive.namelist()) == {"manifest.json", "document.md"}
    assert "<script>" not in markdown
    assert "data:image/svg+xml" not in markdown
    assert "unsupported_embedded_image:image/svg+xml" in manifest["warnings"]
    assert "missing_image_references:1" in manifest["warnings"]
